import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_COLOR = RGBColor(158, 26, 26)    # APSRTC Crimson / #9E1A1A
SECONDARY_COLOR = RGBColor(26, 54, 93)  # Navy #1A365D
DARK_TEXT = RGBColor(33, 37, 41)        # #212529
MUTED_TEXT = RGBColor(108, 117, 125)    # #6C757D
CODE_BG = "F8F9FA"

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header(doc, title, subtitle, member_name, member_roll, member_role):
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("AMRITA VISHWA VIDYAPEETHAM\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    r_inst.bold = True
    r_inst.font.size = Pt(11)
    r_inst.font.color.rgb = SECONDARY_COLOR

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj = p_proj.add_run("PROJECT: AMMA TRAVELS (APSRTC SMART BUS BOOKING SYSTEM)")
    r_proj.bold = True
    r_proj.font.size = Pt(14)
    r_proj.font.color.rgb = PRIMARY_COLOR

    p_ded = doc.add_paragraph()
    p_ded.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ded = p_ded.add_run("Dedicated with reverence to Sri Mata Amritanandamayi Devi (Amma)\nFounder Chancellor, Amrita Vishwa Vidyapeetham")
    r_ded.italic = True
    r_ded.font.size = Pt(9.5)
    r_ded.font.color.rgb = MUTED_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    data = [
        ("Document Title:", title),
        ("Student Name & Roll No:", f"{member_name} ({member_roll})"),
        ("Assigned Module & Role:", member_role),
        ("System Architecture Scope:", subtitle)
    ]

    for idx, (lbl, val) in enumerate(data):
        row = tbl.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.5)

        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(lbl)
        r1.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = SECONDARY_COLOR

        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.bold = (idx == 0 or idx == 1 or idx == 2)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.color.rgb = PRIMARY_COLOR
        run.font.size = Pt(13)
        run.bold = True
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.color.rgb = SECONDARY_COLOR
        run.font.size = Pt(11)
        run.bold = True
    return h

def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = DARK_TEXT
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = DARK_TEXT
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = DARK_TEXT
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = DARK_TEXT
    return p

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        r_cap = p_cap.add_run(f"Code Implementation: {caption}")
        r_cap.italic = True
        r_cap.bold = True
        r_cap.font.size = Pt(8.5)
        r_cap.font.color.rgb = SECONDARY_COLOR
        p_cap.paragraph_format.space_after = Pt(2)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_background(cell, CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_callout(doc, text, title="MODULE HIGHLIGHT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_background(cell, "FFF8E7")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r_tit = p.add_run(f"[{title}]\n")
    r_tit.bold = True
    r_tit.font.size = Pt(9)
    r_tit.font.color.rgb = PRIMARY_COLOR
    r_body = p.add_run(text)
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_table_data(doc, headers, rows_data):
    tbl = doc.add_table(rows=len(rows_data)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_background(hdr_cells[i], "9E1A1A")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=110, right=110)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, r_data in enumerate(rows_data):
        row_cells = tbl.rows[r_idx+1].cells
        bg_color = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=110, right=110)
            p = row_cells[c_idx].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_image_figure(doc, img_path, caption_text, width=Inches(5.2)):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Figure Demonstration: {caption_text}")
        r_cap.italic = True
        r_cap.bold = True
        r_cap.font.size = Pt(8.5)
        r_cap.font.color.rgb = SECONDARY_COLOR

print("Person doc utilities loaded.")
