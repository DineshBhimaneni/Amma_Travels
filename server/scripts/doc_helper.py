import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

PRIMARY_COLOR = RGBColor(158, 26, 26)    # APSRTC Crimson / #9E1A1A
SECONDARY_COLOR = RGBColor(26, 54, 93)  # Navy Blue / #1A365D
DARK_TEXT = RGBColor(33, 37, 41)        # #212529
MUTED_TEXT = RGBColor(108, 117, 125)    # #6C757D
CODE_BG = "F8F9FA"                      # Light Grey
HEADER_BG = "9E1A1A"                    # Table Header Crimson

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header(doc, title, subtitle, member_name, member_role, member_roll="AM.EN.U4CSEXXXXX"):
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("AMRITA VISHWA VIDYAPEETHAM\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    r_inst.bold = True
    r_inst.font.size = Pt(11)
    r_inst.font.color.rgb = SECONDARY_COLOR

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj = p_proj.add_run("PROJECT: AMMA TRAVELS (APSRTC SMART BUS BOOKING SYSTEM)")
    r_proj.bold = True
    r_proj.font.size = Pt(14)
    r_proj.font.color.rgb = PRIMARY_COLOR

    p_ded = doc.add_paragraph()
    p_ded.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ded = p_ded.add_run("Dedicated with deep gratitude and reverence to Sri Mata Amritanandamayi Devi (Amma)\nFounder & Chancellor, Amrita Vishwa Vidyapeetham")
    r_ded.italic = True
    r_ded.font.size = Pt(9.5)
    r_ded.font.color.rgb = MUTED_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Info Table
    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    data = [
        ("Document Title:", title),
        ("Assigned Module & Role:", f"{member_role}"),
        ("Team Member Name:", f"{member_name} (Roll No: {member_roll})"),
        ("System Component:", subtitle)
    ]

    for idx, (lbl, val) in enumerate(data):
        row = tbl.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.5)

        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(lbl)
        r1.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = SECONDARY_COLOR

        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.bold = (idx == 0 or idx == 1)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.color.rgb = PRIMARY_COLOR
        run.font.size = Pt(13)
        run.bold = True
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.color.rgb = SECONDARY_COLOR
        run.font.size = Pt(11)
        run.bold = True
    return h

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        r_cap = p_cap.add_run(f"Code Listing: {caption}")
        r_cap.italic = True
        r_cap.bold = True
        r_cap.font.size = Pt(8.5)
        r_cap.font.color.rgb = SECONDARY_COLOR
        p_cap.paragraph_format.space_after = Pt(2)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_background(cell, CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_callout(doc, text, title="NOTE / ARCHITECTURE HIGHLIGHT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_background(cell, "FFF8E7") # Light Gold Tint
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r_tit = p.add_run(f"[{title}]\n")
    r_tit.bold = True
    r_tit.font.size = Pt(9)
    r_tit.font.color.rgb = PRIMARY_COLOR
    r_body = p.add_run(text)
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Base setup ready.")
